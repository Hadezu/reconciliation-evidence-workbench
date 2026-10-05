import io
import json
import subprocess
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from reconproof import ingest, pipeline
from reconproof.engine import calculate
from reconproof.export import workbook
from reconproof.rules import load_rules

ROOT = Path(__file__).resolve().parents[1]


def rules():
    return load_rules((ROOT / "examples/rules.json").read_bytes())


def csv(rows):
    return ("invoice,currency,gross,issued_at,state\n" + "\n".join(rows)).encode()


def test_empty_inputs_not_false_pass():
    report = calculate([], rules(), ":memory:")
    assert report["verdict"] == "NO_COMPARABLE_DATA"


def test_only_excluded_not_false_pass():
    rows = ingest.normalize(
        csv(["A,EUR,1.00,2020-01-01,posted"]), ".csv", "left", rules()
    )
    assert calculate(rows, rules(), ":memory:")["verdict"] == "NO_COMPARABLE_DATA"


def test_missing_identity_and_unknown_currency_are_visible():
    rows = ingest.normalize(
        csv([",EUR,1.00,2026-09-01,posted", "A,XYZ,2.00,2026-09-01,posted"]),
        ".csv",
        "left",
        rules(),
    )
    result = calculate(rows, rules(), ":memory:")
    assert result["row_accounting"][0]["invalid"] == 2
    assert result["verdict"] == "REVIEW_REQUIRED"
    assert sum(x["rows"] for x in result["totals"]) == 2
    assert (
        next(x for x in result["totals"] if x["currency"] == "XYZ")["known_minor"]
        == 200
    )


def test_record_and_cell_limits(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="records"):
        ingest.normalize(
            csv(["A,EUR,1,2026-09-01,posted"] * 2), ".csv", "left", rules()
        )
    with pytest.raises(ValueError, match="cell"):
        ingest.normalize(
            csv(["A" * 2001 + ",EUR,1,2026-09-01,posted"]), ".csv", "left", rules()
        )


def test_xlsx_date_label_vs_hidden_time():
    book = Workbook()
    sheet = book.active
    sheet.title = "Export"
    sheet.append(["invoice", "currency", "gross", "issued_at", "state"])
    for key, dt in [
        ("date", datetime(2026, 9, 2)),
        ("hidden", datetime(2026, 9, 2, 12)),
    ]:
        sheet.append([key, "EUR", "1.00", dt, "posted"])
        sheet.cell(sheet.max_row, 4).number_format = "yyyy-mm-dd"
    buf = io.BytesIO()
    book.save(buf)
    cfg = rules()
    cfg.left.sheet = "Export"
    rows = ingest.normalize(buf.getvalue(), ".xlsx", "left", cfg)
    assert rows[0]["disposition"] == "eligible"
    assert rows[1]["disposition"] == "invalid"


def test_external_link_archive_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/externalLinks/externalLink1.xml", "x")
    with pytest.raises(ValueError, match="external"):
        ingest.read_table(buf.getvalue(), ".xlsx", rules().right)


def test_oversized_expansion_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("large.xml", b" " * 40_000_001)
    with pytest.raises(ValueError, match="expanded"):
        ingest.read_table(buf.getvalue(), ".xlsx", rules().right)


def test_xml_entity_rejected():
    book = Workbook()
    book.active.title = "Export"
    buf = io.BytesIO()
    book.save(buf)
    target = io.BytesIO()
    with zipfile.ZipFile(buf) as old, zipfile.ZipFile(target, "w") as new:
        for name in old.namelist():
            data = old.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = b'<!DOCTYPE a [<!ENTITY x "EXPAND">]><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&x;</t></is></c></row></sheetData></worksheet>'
            new.writestr(name, data)
    from defusedxml.common import DefusedXmlException

    with pytest.raises(ValueError) as caught:
        ingest.read_table(target.getvalue(), ".xlsx", rules().right)
    assert isinstance(caught.value.__cause__, DefusedXmlException)


def test_long_raw_json_not_truncated(tmp_path):
    rows = ingest.normalize(csv(["A,EUR,1,2026-09-01,posted"]), ".csv", "left", rules())
    rows[0]["raw_json"] = json.dumps({f"column{i}": "x" * 2000 for i in range(30)})
    report = calculate(rows, rules(), ":memory:")
    workbook(report, tmp_path / "r.xlsx")
    book = load_workbook(tmp_path / "r.xlsx")
    cells = list(book["Source rows"].iter_rows(min_row=2, values_only=True))[0]
    assert "".join(cells[8:]) == rows[0]["raw_json"]
    book.close()


@pytest.mark.parametrize("mutation", ["add", "remove"])
def test_manifest_file_set(tmp_path, mutation):
    out = tmp_path / "run"
    pipeline.run(
        ROOT / "examples/ledger.csv",
        ROOT / "examples/target.xlsx",
        ROOT / "examples/rules.json",
        out,
    )
    if mutation == "add":
        (out / "unexpected.txt").write_text("not in the evidence set")
    else:
        (out / "report.html").unlink()
    with pytest.raises(ValueError, match="Missing or unexpected"):
        pipeline.verify(out)


def test_cancelled_bad_amount_is_not_silently_excluded():
    rows = ingest.normalize(
        csv(["A,EUR,bad,2026-09-01,cancelled"]), ".csv", "left", rules()
    )
    assert rows[0]["disposition"] == "invalid"


def test_missing_status_is_invalid_not_an_exclusion():
    rows = ingest.normalize(csv(["A,EUR,1.00,2026-09-01,"]), ".csv", "left", rules())
    assert rows[0]["disposition"] == "invalid"
    assert rows[0]["reason"] == "MISSING_STATUS"


def test_duplicate_groups_preserve_very_large_exact_sums():
    rows = ingest.normalize(
        csv(["A,EUR,999999999999.99,2026-09-01,posted"] * 101), ".csv", "left", rules()
    )
    report = calculate(rows, rules(), ":memory:")
    assert report["results"][0]["left_minor"] == 101 * 99999999999999
    assert report["results"][0]["status"] == "AMBIGUOUS_DUPLICATE"


def test_native_xlsx_number_ignores_text_separator():
    assert ingest.money(123.45, ",") == 12345


def test_large_lineage_has_explicit_workbook_pointer(tmp_path):
    rows = ingest.normalize(
        csv(["A,EUR,1.00,2026-09-01,posted"]), ".csv", "left", rules()
    )
    report = calculate(rows, rules(), ":memory:")
    report["results"][0]["left_rows"] = list(range(2, 25_002))
    workbook(report, tmp_path / "large-lineage.xlsx")
    book = load_workbook(tmp_path / "large-lineage.xlsx")
    assert book["Reconciliation"]["G2"].value.startswith("25000 entries; full lineage")
    book.close()


def test_process_death_never_publishes_partial_report(tmp_path):
    code = """import os,sys
from pathlib import Path
from reconproof import pipeline
pipeline.workbook=lambda *args: os._exit(71)
pipeline.run(*[Path(p) for p in sys.argv[1:]])
"""
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            code,
            str(ROOT / "examples/ledger.csv"),
            str(ROOT / "examples/target.xlsx"),
            str(ROOT / "examples/rules.json"),
            str(tmp_path / "run"),
        ],
        capture_output=True,
    )
    assert result.returncode == 71
    assert not (tmp_path / "run").exists()
    assert len(list(tmp_path.glob(".recon-staging-*"))) == 1


def test_concurrent_publication_has_one_winner(tmp_path):
    target = tmp_path / "run"

    def attempt():
        try:
            return pipeline.run(
                ROOT / "examples/ledger.csv",
                ROOT / "examples/target.xlsx",
                ROOT / "examples/rules.json",
                target,
            )["run_key"]
        except ValueError as error:
            assert "exists" in str(error) or "another process" in str(error)
            return None
        except OSError as error:
            # POSIX ENOTEMPTY / Windows EEXIST from the final no-replacement rename.
            import errno

            assert error.errno in {errno.ENOTEMPTY, errno.EEXIST, errno.EACCES}
            return None

    with ThreadPoolExecutor(max_workers=2) as pool:
        result = list(pool.map(lambda _: attempt(), range(2)))
    assert sum(value is not None for value in result) == 1
    assert pipeline.verify(target) == next(value for value in result if value)
