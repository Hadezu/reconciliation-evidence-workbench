import csv
import io
import json
import subprocess
import sys
from pathlib import Path

import duckdb
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from reconproof import pipeline
from reconproof.engine import calculate
from reconproof.export import html_report, portable, workbook
from reconproof.ingest import money, normalize, read_table, reporting_date
from reconproof.rules import load_rules

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def rules():
    return load_rules((ROOT / "examples/rules.json").read_bytes())


def csv_data(rows):
    target = io.StringIO(newline="")
    writer = csv.writer(target)
    writer.writerow(["invoice", "currency", "gross", "issued_at", "state"])
    writer.writerows(rows)
    return target.getvalue().encode()


def source(
    key="001", amount="10.00", currency="EUR", day="2026-09-10", status="posted"
):
    return [key, currency, amount, day, status]


def inputs(left, right, rules):
    same = rules.model_copy(update={"right": rules.left})
    return normalize(csv_data(left), ".csv", "left", same) + normalize(
        csv_data(right), ".csv", "right", same
    )


def xlsx(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Export"
    for row in [["invoice", "currency", "gross", "issued_at", "state"], *rows]:
        sheet.append(row)
    target = io.BytesIO()
    book.save(target)
    return target.getvalue()


def test_hand_computed_oracle(tmp_path):
    expected = json.loads((ROOT / "examples/expected.json").read_text())
    report = pipeline.run(
        ROOT / "examples/ledger.csv",
        ROOT / "examples/target.xlsx",
        ROOT / "examples/rules.json",
        tmp_path / "run",
    )
    assert report["verdict"] == expected["verdict"]
    assert len(report["rows"]) == expected["source_rows"]
    assert report["exception_groups"] == expected["exception_groups"]
    assert {
        r["key"] + "|" + r["currency"]: r["status"] for r in report["results"]
    } == expected["groups"]
    for side in ["left", "right"]:
        counts = next(x for x in report["row_accounting"] if x["side"] == side)
        assert {k: counts[k] for k in ["eligible", "invalid", "excluded"]} == expected[
            side + "_accounting"
        ]
    euro = next(b for b in report["bridges"] if b["currency"] == "EUR")
    for key in ["left_eligible_minor", "right_eligible_minor", "difference_minor"]:
        assert euro[key] == expected["eur_" + key]
    assert pipeline.verify(tmp_path / "run") == report["run_key"]
    with duckdb.connect(str(tmp_path / "run/evidence.duckdb"), read_only=True) as db:
        assert db.execute("select count(*) from source_rows").fetchone()[0] == 25
    exported = json.loads((tmp_path / "run/report.json").read_text())
    assert isinstance(exported["bridges"][0]["left_eligible_minor"], str)
    assert exported["rules"] == report["rules"]


@pytest.mark.parametrize(
    "value",
    [
        "NaN",
        "inf",
        "1e3",
        "1,000.00",
        "12.345",
        "",
        None,
        True,
        "1000000000000",
        "--1",
        "=1+2",
    ],
)
def test_invalid_money(value):
    with pytest.raises(ValueError):
        money(value, ".")


@pytest.mark.parametrize(
    "value,expected",
    [
        ("0.10", 10),
        ("-12.01", -1201),
        ("+1", 100),
        ("000.09", 9),
        ("999999999999.99", 99999999999999),
    ],
)
def test_exact_money(value, expected):
    assert money(value, ".") == expected


def test_explicit_comma_decimal():
    assert money("12,34", ",") == 1234
    with pytest.raises(ValueError):
        money("12.34", ",")


@pytest.mark.parametrize(
    "value",
    ["2026-02-30", "09/10/2026", "2026-09-10T12:00:00", "20260910", "yesterday"],
)
def test_invalid_dates(value):
    with pytest.raises(ValueError):
        reporting_date(value)


def test_timezone_cutover():
    assert reporting_date("2026-10-01T00:30:00+02:00") == "2026-09-30"
    assert reporting_date("2026-09-01T00:30:00+02:00") == "2026-08-31"


@pytest.mark.parametrize(
    "left,right,status",
    [
        ([source()], [source()], "MATCH"),
        ([source(), source()], [source()], "AMBIGUOUS_DUPLICATE"),
        ([source()], [source(), source()], "AMBIGUOUS_DUPLICATE"),
        ([source(amount="bad")], [source()], "INVALID_GROUP"),
        ([source(amount="9.99")], [source()], "AMOUNT_MISMATCH"),
        ([source()], [], "LEFT_ONLY"),
        ([], [source()], "RIGHT_ONLY"),
    ],
)
def test_classification(tmp_path, rules, left, right, status):
    result = calculate(inputs(left, right, rules), rules, tmp_path / "run.duckdb")
    assert result["results"][0]["status"] == status


def test_invalid_duplicate_cannot_disappear(tmp_path, rules):
    rows = inputs([source(), source(amount="bad")], [source()], rules)
    report = calculate(rows, rules, tmp_path / "data.duckdb")
    assert report["results"][0]["status"] == "AMBIGUOUS_DUPLICATE"
    assert report["results"][0]["left_rows"] == [2, 3]


def test_currency_never_cross_matches(tmp_path, rules):
    report = calculate(
        inputs([source(currency="EUR")], [source(currency="PLN")], rules),
        rules,
        tmp_path / "d.duckdb",
    )
    assert {r["status"] for r in report["results"]} == {"LEFT_ONLY", "RIGHT_ONLY"}


def test_tolerance_is_explicit_nonzero_delta(tmp_path, rules):
    rules = rules.model_copy(update={"tolerance_minor": 1})
    report = calculate(
        inputs([source(amount="0.10")], [source(amount="0.11")], rules),
        rules,
        tmp_path / "d.duckdb",
    )
    assert report["results"][0]["status"] == "WITHIN_TOLERANCE"
    assert report["bridges"][0]["difference_minor"] == -1
    assert report["verdict"] == "RECONCILED"


def test_blank_extra_and_short_rows_preserved(rules):
    data = b"invoice,currency,gross,issued_at,state\n\nA,EUR,1.00,2026-09-01,posted,extra\nB,EUR\n"
    rows = normalize(data, ".csv", "left", rules)
    assert len(rows) == 3
    assert all(r["disposition"] == "invalid" for r in rows)
    assert "extra" in rows[1]["raw_json"]


@pytest.mark.parametrize(
    "data", [b"", b"a,a\n1,2", b"x,y\n1,2", b'"unterminated', b"\xff"]
)
def test_bad_files_fail_whole_run(data, rules):
    with pytest.raises((ValueError, UnicodeError, csv.Error)):
        read_table(data, ".csv", rules.left)


def test_numeric_xlsx_id_and_formula_rejected(rules):
    rules = rules.model_copy(
        update={"left": rules.left.model_copy(update={"sheet": "Export"})}
    )
    data = xlsx([source(key=17), source(key="00017"), source(key="F", amount="=1+2")])
    rows = normalize(data, ".xlsx", "left", rules)
    assert rows[0]["disposition"] == "invalid"
    assert rows[1]["key"] == "00017" and rows[1]["disposition"] == "eligible"
    assert "FORMULA_OR_ERROR_CELL" in rows[2]["reason"]


def test_xlsx_explicit_sheet(rules):
    with pytest.raises(ValueError, match="sheet"):
        read_table(xlsx([source()]), ".xlsx", rules.left)


def test_malicious_formula_text_export(tmp_path, rules):
    key = '=HYPERLINK("https://example.invalid","x")'
    report = calculate(
        inputs([source(key=key)], [source(key=key)], rules),
        rules,
        tmp_path / "d.duckdb",
    )
    workbook(report, tmp_path / "report.xlsx")
    book = load_workbook(tmp_path / "report.xlsx", data_only=False)
    assert book["Reconciliation"]["A2"].value == key
    assert book["Reconciliation"]["A2"].data_type == "s"
    assert not any(c.data_type == "f" for s in book for row in s for c in row)
    book.close()


def test_html_injection_and_large_integer(tmp_path, rules):
    key = "</script><script>alert(1)</script>"
    report = calculate(
        inputs([source(key=key)], [source(key=key)], rules),
        rules,
        tmp_path / "d.duckdb",
    )
    report["provenance"] = {}
    report["run_key"] = "test"
    html_report(report, tmp_path / "report.html")
    assert key not in (tmp_path / "report.html").read_text()
    assert (
        portable({"known_minor": 999999999999999999})["known_minor"]
        == "999999999999999999"
    )


def test_repeated_run_stable_inputs_and_report(tmp_path):
    paths = [
        ROOT / "examples/ledger.csv",
        ROOT / "examples/target.xlsx",
        ROOT / "examples/rules.json",
    ]
    before = [p.read_bytes() for p in paths]
    first = pipeline.run(*paths, tmp_path / "first")
    second = pipeline.run(*paths, tmp_path / "second")
    assert first == second
    assert (tmp_path / "first/report.json").read_bytes() == (
        tmp_path / "second/report.json"
    ).read_bytes()
    assert before == [p.read_bytes() for p in paths]
    with pytest.raises(ValueError, match="exists"):
        pipeline.run(*paths, tmp_path / "first")


def test_failed_export_leaves_no_published_run(tmp_path, monkeypatch):
    def fail(*_):
        raise OSError("simulated disk failure")

    monkeypatch.setattr(pipeline, "workbook", fail)
    with pytest.raises(OSError):
        pipeline.run(
            ROOT / "examples/ledger.csv",
            ROOT / "examples/target.xlsx",
            ROOT / "examples/rules.json",
            tmp_path / "run",
        )
    assert not (tmp_path / "run").exists()
    assert not list(tmp_path.glob(".recon-staging-*"))


def test_tamper_detection(tmp_path):
    pipeline.run(
        ROOT / "examples/ledger.csv",
        ROOT / "examples/target.xlsx",
        ROOT / "examples/rules.json",
        tmp_path / "run",
    )
    (tmp_path / "run/report.json").write_text("altered")
    with pytest.raises(ValueError, match="integrity"):
        pipeline.verify(tmp_path / "run")


def test_unknown_rules_and_duplicate_keys():
    raw = (ROOT / "examples/rules.json").read_bytes()
    with pytest.raises(ValueError):
        load_rules(raw.replace(b'"schema_version": 1', b'"typo": 1'))
    with pytest.raises(ValueError, match="Duplicate JSON"):
        load_rules(b'{"a":1,"a":2}')


def test_cli_review_exit_and_verify(tmp_path):
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "reconproof.cli",
            "run",
            "--left",
            str(ROOT / "examples/ledger.csv"),
            "--right",
            str(ROOT / "examples/target.xlsx"),
            "--rules",
            str(ROOT / "examples/rules.json"),
            "--out",
            str(tmp_path / "run"),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 2 and "REVIEW_REQUIRED" in result.stdout
    verified = subprocess.run(
        [sys.executable, "-m", "reconproof.cli", "verify", str(tmp_path / "run")],
        capture_output=True,
        text=True,
    )
    assert verified.returncode == 0 and "INTEGRITY_OK" in verified.stdout


@given(
    st.lists(
        st.integers(min_value=-99999999, max_value=99999999), min_size=1, max_size=30
    )
)
@settings(max_examples=35, deadline=None)
def test_permutation_and_exact_sum_property(amounts):
    rules = load_rules((ROOT / "examples/rules.json").read_bytes())
    records = [
        source(
            key=f"I{i}",
            amount=f"{'-' if n < 0 else ''}{abs(n) // 100}.{abs(n) % 100:02d}",
        )
        for i, n in enumerate(amounts)
    ]
    rows = inputs(records, list(reversed(records)), rules)
    report = calculate(rows, rules, ":memory:")
    assert report["verdict"] == "RECONCILED"
    assert all(r["status"] == "MATCH" for r in report["results"])
    assert report["bridges"][0]["left_eligible_minor"] == sum(amounts)
