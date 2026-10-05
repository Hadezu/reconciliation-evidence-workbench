"""Static handover files. Input text is never emitted as a spreadsheet formula or HTML markup."""

import hashlib
import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def dump_json(value):
    return (
        json.dumps(value, ensure_ascii=True, sort_keys=True, indent=2, allow_nan=False)
        + "\n"
    )


def sha(data):
    return hashlib.sha256(data).hexdigest()


def text_amount(minor):
    return f"{'-' if minor < 0 else ''}{abs(minor) // 100}.{abs(minor) % 100:02d}"


def portable(value):
    """Keep large exact money values safe for JavaScript/JSON consumers."""
    if isinstance(value, dict):
        return {
            k: str(v)
            if k.endswith("_minor") and k != "tolerance_minor" and isinstance(v, int)
            else portable(v)
            for k, v in value.items()
        }
    if isinstance(value, list):
        return [portable(x) for x in value]
    return value


def excel_safe(value):
    # Preserve XML-illegal characters as visible escape sequences, rather than drop data.
    import re

    return re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda m: f"\\u{ord(m[0]):04x}", value
    )


def workbook(report, path):
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            for column, value in enumerate(row, 1):
                cell = ws.cell(ws.max_row + 1 if column == 1 else ws.max_row, column)
                if value is None:
                    cell.value = ""
                elif isinstance(value, (dict, list)):
                    encoded = json.dumps(value, ensure_ascii=True)
                    if len(encoded) > 30_000:
                        encoded = f"{len(value)} entries; full lineage in report.json and Source rows sheet (filter reference + currency)."
                    cell.value, cell.data_type = encoded, "s"
                elif isinstance(value, str):
                    # Explicit string cells preserve leading zeros and disable formula injection.
                    cell.value, cell.data_type = excel_safe(value), "s"
                else:
                    cell.value = value
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183B4E")
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24 if col < 8 else 40
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        return ws

    summary = [
        ["Verdict", report["verdict"]],
        ["Case", "Local reconciliation under explicit rules; not an accounting audit"],
        [
            "Money",
            "Exact decimal text; no formulas or FX. Known subtotals do not include unparseable amounts.",
        ],
        ["Tolerance (minor units)", report["rules"]["tolerance_minor"]],
    ]
    for bridge in report["bridges"]:
        summary += [
            [bridge["currency"] + " / " + key, text_amount(bridge[key])]
            for key in (
                "left_eligible_minor",
                "right_eligible_minor",
                "difference_minor",
            )
        ]
    for row in report["row_accounting"]:
        summary += [
            [row["side"] + " / " + key, row[key]]
            for key in ("source_rows", "eligible", "invalid", "excluded")
        ]
    sheet("Summary", ["Measure", "Value"], summary)
    sheet(
        "Reconciliation",
        [
            "Reference",
            "Currency",
            "Status",
            "Left known amount",
            "Right known amount",
            "Known delta",
            "Left source rows",
            "Right source rows",
        ],
        [
            [
                r["key"],
                r["currency"],
                r["status"],
                text_amount(r["left_minor"]),
                text_amount(r["right_minor"]),
                text_amount(r["known_delta_minor"]),
                r["left_rows"],
                r["right_rows"],
            ]
            for r in report["results"]
        ],
    )
    # Split long JSON across cells; Excel truncates cell text beyond 32,767 characters.
    chunks = max(
        [1] + [(len(r["raw_json"]) + 29_999) // 30_000 for r in report["rows"]]
    )
    sheet(
        "Source rows",
        [
            "Source",
            "Record / worksheet row",
            "Reference",
            "Currency",
            "Parsed amount",
            "UTC reporting date",
            "Disposition",
            "Reason",
        ]
        + [f"Original JSON part {i + 1}" for i in range(chunks)],
        [
            [
                r["side"],
                r["source_row"],
                r["key"],
                r["currency"],
                text_amount(r["amount_minor"])
                if r["amount_minor"] is not None
                else "UNPARSEABLE",
                r["report_date"],
                r["disposition"],
                r["reason"],
            ]
            + [r["raw_json"][i * 30_000 : (i + 1) * 30_000] for i in range(chunks)]
            for r in report["rows"]
        ],
    )
    sheet(
        "Rules",
        ["Rule", "Value"],
        [[key, value] for key, value in report["rules"].items()],
    )
    book.save(path)


def html_report(report, path):
    template = Path(__file__).with_name("report.html").read_text(encoding="utf-8")
    # An application/json script element still needs '<' escaped to prevent closing-tag injection.
    payload = (
        dump_json(portable(report))
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace("&", "\\u0026")
    )
    output = template.replace("__REPORT_DATA__", payload).replace(
        "__VERDICT__", html.escape(report["verdict"])
    )
    path.write_text(output, encoding="utf-8")
