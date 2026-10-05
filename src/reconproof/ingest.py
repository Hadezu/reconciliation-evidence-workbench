"""Bounded file readers. Preserve every tabular record, including blank/invalid rows."""

import csv
import io
import json
import re
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_datetime

from .rules import Rules, SourceRule

MAX_BYTES = 8_000_000
MAX_ROWS = 25_000
MAX_COLUMNS = 40
MAX_CELL = 2_000


def primitive(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def checked_values(values):
    if len(values) > MAX_COLUMNS or any(len(str(x or "")) > MAX_CELL for x in values):
        raise ValueError("Column/cell size limit exceeded")
    return values


def cell_value(cell):
    value = cell.value
    if isinstance(value, datetime) and is_datetime(cell.number_format) == "date":
        # Only date-formatted midnight is a date label. Never hide a nonzero time.
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
    return primitive(value)


def read_table(data: bytes, suffix: str, rule: SourceRule):
    if len(data) > MAX_BYTES:
        raise ValueError("Input exceeds 8 MB")
    if suffix == ".csv":
        reader = csv.reader(
            io.StringIO(data.decode("utf-8-sig"), newline=""),
            delimiter=rule.delimiter,
            strict=True,
        )
        table = []
        for row in reader:
            if len(table) > MAX_ROWS:
                raise ValueError("Input exceeds 25,000 records")
            table.append((checked_values(row), set()))
    elif suffix == ".xlsx":
        # Reject oversized archives before openpyxl expands XML/shared strings.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(x.file_size for x in members) > 40_000_000:
                raise ValueError("XLSX expanded size exceeds limit")
            if any(x.flag_bits & 1 for x in members):
                raise ValueError("Encrypted XLSX unsupported")
            if any(
                "vbaProject" in x.filename or x.filename.startswith("xl/externalLinks/")
                for x in members
            ):
                raise ValueError("Macros/external workbook links unsupported")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if not rule.sheet or rule.sheet not in workbook.sheetnames:
                raise ValueError("XLSX requires an explicit existing sheet name")
            sheet = workbook[rule.sheet]
            sheet.reset_dimensions()  # Do not trust producer-supplied used-range metadata.
            table = []
            for cells in sheet.iter_rows():
                if len(table) > MAX_ROWS:
                    raise ValueError("Input exceeds 25,000 records")
                table.append(
                    (
                        checked_values([cell_value(c) for c in cells]),
                        {i for i, c in enumerate(cells) if c.data_type in {"f", "e"}},
                    )
                )
        finally:
            workbook.close()
    else:
        raise ValueError("Only .csv and .xlsx inputs supported")
    if not table:
        raise ValueError("Missing header")
    headers, header_errors = table[0]
    if (
        header_errors
        or not headers
        or len(headers) > MAX_COLUMNS
        or any(not isinstance(x, str) or not x.strip() for x in headers)
        or len(headers) != len(set(headers))
    ):
        raise ValueError("Headers must be distinct nonempty text, at most 40 columns")
    if not set(rule.columns.values()) <= set(headers):
        raise ValueError("Mapped column missing")
    for values, _ in table:
        if len(values) > MAX_COLUMNS or any(
            len(str(x or "")) > MAX_CELL for x in values
        ):
            raise ValueError("Column/cell size limit exceeded")
    return headers, table[1:]


def money(value, separator):
    if value is None or isinstance(value, bool):
        raise ValueError("INVALID_AMOUNT")
    if isinstance(value, (int, float)):
        separator = "."  # Native numeric XLSX cells are not localized text.
    token = str(value).strip()
    pattern = r"[+-]?\d{1,12}(?:" + re.escape(separator) + r"\d{1,2})?"
    if not re.fullmatch(pattern, token):
        raise ValueError("INVALID_AMOUNT")
    # No float arithmetic and no silent rounding of additional decimal places.
    return int(Decimal(token.replace(separator, ".")) * 100)


def reporting_date(value):
    if not isinstance(value, str):
        raise ValueError("INVALID_DATE")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return date.fromisoformat(value).isoformat()
    timestamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if timestamp.tzinfo is None or "T" not in value:
        raise ValueError("AMBIGUOUS_DATE_TIME")
    return timestamp.astimezone(UTC).date().isoformat()


def normalize(data: bytes, suffix: str, side: str, rules: Rules):
    rule = getattr(rules, side)
    headers, records = read_table(data, suffix, rule)
    rows = []
    for ordinal, (values, invalid_cells) in enumerate(records, 2):
        raw = {
            header: values[i] if i < len(values) else None
            for i, header in enumerate(headers)
        }
        if len(values) > len(headers):
            raw["__extra_cells__"] = values[len(headers) :]
        errors = []
        if len(values) != len(headers):
            errors.append("COLUMN_COUNT")
        mapped = {field: raw[column] for field, column in rule.columns.items()}
        if invalid_cells:
            errors.append("FORMULA_OR_ERROR_CELL")
        key = mapped["key"].strip() if isinstance(mapped["key"], str) else ""
        if not key:
            errors.append("KEY_MUST_BE_NONEMPTY_TEXT")
        currency = str(mapped["currency"] or "").strip().upper()
        if currency not in rules.currencies:
            errors.append("UNSUPPORTED_CURRENCY")
        amount, day = None, None
        try:
            amount = money(mapped["amount"], rule.decimal_separator)
        except ValueError as error:
            errors.append(str(error))
        try:
            day = reporting_date(mapped["date"])
        except (ValueError, TypeError, OverflowError):
            errors.append("INVALID_OR_AMBIGUOUS_DATE")
        status = str(mapped["status"] if mapped["status"] is not None else "").strip()
        if not status:
            errors.append("MISSING_STATUS")
        disposition, reason = "eligible", "IN_SCOPE"
        if errors:
            disposition, reason = "invalid", ";".join(errors)
        elif status not in rule.included_statuses:
            disposition, reason = "excluded", "STATUS_OUTSIDE_RULE"
        elif not rules.period_start <= day < rules.period_end_exclusive:
            disposition, reason = "excluded", "OUTSIDE_PERIOD"
        rows.append(
            {
                "side": side,
                "source_row": ordinal,
                "key": key,
                "currency": currency,
                "amount_minor": amount,
                "report_date": day,
                "disposition": disposition,
                "reason": reason,
                "raw_json": json.dumps(
                    raw, ensure_ascii=True, sort_keys=True, allow_nan=False
                ),
            }
        )
    return rows
